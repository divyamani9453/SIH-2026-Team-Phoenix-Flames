import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Colors
    NAVY = RGBColor(26, 54, 93)
    BLUE = RGBColor(13, 110, 253)
    BANNER_BLUE = RGBColor(0, 114, 188)
    DARK_TEXT = RGBColor(30, 41, 59)
    LIGHT_BG = RGBColor(248, 250, 252)
    MUTED_TEXT = RGBColor(100, 116, 139)
    WHITE = RGBColor(255, 255, 255)
    BOX_BG = RGBColor(241, 245, 249)
    BORDER_COLOR = RGBColor(203, 213, 225)

    def add_common_header_footer(slide, title_text, slide_num):
        # Team Name Oval (Top Left)
        shape_oval = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(0.5), Inches(0.3), Inches(1.4), Inches(0.8)
        )
        shape_oval.fill.background()
        shape_oval.line.color.rgb = RGBColor(147, 112, 219)
        shape_oval.line.width = Pt(1.5)
        tf_oval = shape_oval.text_frame
        tf_oval.word_wrap = True
        p_oval = tf_oval.paragraphs[0]
        p_oval.text = "Your Team Name"
        p_oval.font.size = Pt(11)
        p_oval.font.bold = True
        p_oval.font.color.rgb = NAVY
        p_oval.alignment = PP_ALIGN.CENTER

        # Title (Center Top)
        txBox = slide.shapes.add_textbox(Inches(2.2), Inches(0.3), Inches(8.0), Inches(0.8))
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(26)
        p.font.bold = True
        p.font.color.rgb = RGBColor(15, 23, 42)
        p.alignment = PP_ALIGN.CENTER

        # SIH Header Text/Logo (Top Right)
        tx_sih = slide.shapes.add_textbox(Inches(10.3), Inches(0.2), Inches(2.5), Inches(0.9))
        tf_sih = tx_sih.text_frame
        p_sih1 = tf_sih.paragraphs[0]
        p_sih1.text = "SMART INDIA"
        p_sih1.font.size = Pt(13)
        p_sih1.font.bold = True
        p_sih1.font.color.rgb = NAVY
        p_sih1.alignment = PP_ALIGN.RIGHT
        p_sih2 = tf_sih.add_paragraph()
        p_sih2.text = "HACKATHON 2026"
        p_sih2.font.size = Pt(13)
        p_sih2.font.bold = True
        p_sih2.font.color.rgb = NAVY
        p_sih2.alignment = PP_ALIGN.RIGHT

        # Bottom Banner Bar
        banner = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.0), Inches(13.333), Inches(0.5)
        )
        banner.fill.solid()
        banner.fill.fore_color.rgb = BANNER_BLUE
        banner.line.color.rgb = BANNER_BLUE

        # Banner Footer Text
        tx_foot = slide.shapes.add_textbox(Inches(0.5), Inches(7.05), Inches(12.333), Inches(0.4))
        tf_foot = tx_foot.text_frame
        p_foot = tf_foot.paragraphs[0]
        p_foot.text = "@SIH Idea submission- Template"
        p_foot.font.size = Pt(11)
        p_foot.font.color.rgb = WHITE

        p_num = tf_foot.add_paragraph()
        p_num.text = str(slide_num)
        p_num.font.size = Pt(11)
        p_num.font.bold = True
        p_num.font.color.rgb = WHITE
        p_num.alignment = PP_ALIGN.RIGHT

    blank_layout = prs.slide_layouts[6]

    # -------------------------------------------------------------
    # SLIDE 1: TITLE PAGE
    # -------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)

    # Title Header Banner
    tx_title_hdr = slide1.shapes.add_textbox(Inches(1.0), Inches(0.4), Inches(11.333), Inches(0.8))
    p_th = tx_title_hdr.text_frame.paragraphs[0]
    p_th.text = "SMART INDIA HACKATHON 2026"
    p_th.font.size = Pt(32)
    p_th.font.bold = True
    p_th.font.color.rgb = BANNER_BLUE
    p_th.alignment = PP_ALIGN.CENTER

    tx_tp = slide1.shapes.add_textbox(Inches(1.0), Inches(1.2), Inches(11.333), Inches(0.6))
    p_tp = tx_tp.text_frame.paragraphs[0]
    p_tp.text = "TITLE PAGE"
    p_tp.font.size = Pt(28)
    p_tp.font.bold = True
    p_tp.font.color.rgb = NAVY
    p_tp.alignment = PP_ALIGN.CENTER

    # Content Container
    container1 = slide1.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(2.0), Inches(11.333), Inches(4.8)
    )
    container1.fill.solid()
    container1.fill.fore_color.rgb = LIGHT_BG
    container1.line.color.rgb = BORDER_COLOR

    tf1 = container1.text_frame
    tf1.word_wrap = True
    tf1.margin_left = Inches(0.5)
    tf1.margin_top = Inches(0.4)

    items_s1 = [
        ("Problem Statement ID:", " 1812 (PS 083)"),
        ("Problem Statement Title:", " Development of Heatwave Warning and Biometeorological Decision Support System"),
        ("Theme:", " Disaster Management, Climate Resilience & Smart Cities"),
        ("PS Category:", " Software"),
        ("Team ID:", " [Team ID Placeholder - e.g., SIH2026-TEAM-8301]"),
        ("Team Name (Registered on portal):", " [Team Name Placeholder - e.g., HeatResilience AI]")
    ]

    for idx, (label, val) in enumerate(items_s1):
        p = tf1.paragraphs[0] if idx == 0 else tf1.add_paragraph()
        p.space_after = Pt(16)

        run_lbl = p.add_run()
        run_lbl.text = "• " + label
        run_lbl.font.size = Pt(20)
        run_lbl.font.bold = True
        run_lbl.font.color.rgb = NAVY

        run_val = p.add_run()
        run_val.text = val
        run_val.font.size = Pt(20)
        run_val.font.color.rgb = DARK_TEXT

    # -------------------------------------------------------------
    # SLIDE 2: IDEA TITLE - PROPOSED SOLUTION
    # -------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    add_common_header_footer(slide2, "IDEA TITLE", 2)

    # Subtitle
    tx_sub2 = slide2.shapes.add_textbox(Inches(0.5), Inches(1.1), Inches(12.333), Inches(0.5))
    p_s2 = tx_sub2.text_frame.paragraphs[0]
    p_s2.text = "❖ Proposed Solution (Describe your Idea/Solution/Prototype)"
    p_s2.font.size = Pt(22)
    p_s2.font.bold = True
    p_s2.font.color.rgb = BANNER_BLUE

    # 3 Columns Layout
    col_width = Inches(3.9)
    col_gap = Inches(0.3)
    left_start = Inches(0.5)

    s2_sections = [
        ("Detailed Explanation", [
            "Real-time calculations of UTCI, ISO 7243 WBGT, and NOAA Heat Index across 641 districts & 291 municipal wards.",
            "Random Forest ML Heat Vulnerability Index (HVI) engine using canopy cover, albedo, and demographic metrics.",
            "Urban Cooling Policy Simulator predicting temperature drops for green roofs & cool pavements.",
            "Automated multi-channel emergency alert gateway (SMS/WhatsApp)."
        ]),
        ("Addressing the Problem", [
            "Moves beyond simple temperature metrics to evaluate real physiological heat stress (humidity, radiation, wind).",
            "Resolves hyper-local microclimate gaps at municipal ward levels to combat Urban Heat Island (UHI) effects.",
            "Automates IMD/NDMA standard color-coded early warnings (Green, Yellow, Orange, Red) with actionable advisories."
        ]),
        ("Innovation & Uniqueness", [
            "Tri-Index Biometeorological Model executing Fiala 12-node thermoregulation, WBGT, and HI simultaneously.",
            "Physics-informed radiative microclimate simulation for urban intervention impact assessment.",
            "Ultra-lightweight footprint (<180 MB RAM) with SQLite LRU caching (<10k daily weather API credit guard)."
        ])
    ]

    for col_idx, (col_title, col_bullets) in enumerate(s2_sections):
        col_left = left_start + col_idx * (col_width + col_gap)
        card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, col_left, Inches(1.7), col_width, Inches(5.1))
        card.fill.solid()
        card.fill.fore_color.rgb = LIGHT_BG
        card.line.color.rgb = BORDER_COLOR

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.2)
        tf.margin_right = Inches(0.2)
        tf.margin_top = Inches(0.2)

        p_h = tf.paragraphs[0]
        p_h.text = col_title
        p_h.font.size = Pt(16)
        p_h.font.bold = True
        p_h.font.color.rgb = NAVY
        p_h.space_after = Pt(12)

        for b in col_bullets:
            p_b = tf.add_paragraph()
            p_b.text = "• " + b
            p_b.font.size = Pt(13)
            p_b.font.color.rgb = DARK_TEXT
            p_b.space_after = Pt(8)

    # -------------------------------------------------------------
    # SLIDE 3: TECHNICAL APPROACH
    # -------------------------------------------------------------
    slide3 = prs.slides.add_slide(blank_layout)
    add_common_header_footer(slide3, "TECHNICAL APPROACH", 3)

    # Left Box: Tech Stack
    card_tech = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.2), Inches(5.8), Inches(5.6))
    card_tech.fill.solid()
    card_tech.fill.fore_color.rgb = LIGHT_BG
    card_tech.line.color.rgb = BORDER_COLOR
    tf_tech = card_tech.text_frame
    tf_tech.word_wrap = True
    tf_tech.margin_left = Inches(0.3)
    tf_tech.margin_top = Inches(0.3)

    p_th = tf_tech.paragraphs[0]
    p_th.text = "• Technologies to be used"
    p_th.font.size = Pt(18)
    p_th.font.bold = True
    p_th.font.color.rgb = NAVY
    p_th.space_after = Pt(12)

    tech_list = [
        ("Core Stack:", " Python 3.12, Dash by Plotly, Dash Bootstrap Components"),
        ("Biomet & ML:", " Scikit-learn (Random Forest), Pythermalcomfort, NumPy, Pandas"),
        ("Geospatial & Viz:", " GeoPandas, Shapely, Plotly Mapbox Vector Tiles, Openpyxl"),
        ("Data & Caching:", " Open-Meteo Weather API, Requests-Cache (1-hour SQLite LRU)"),
        ("Deployment:", " Gunicorn WSGI, Containerized Docker / Render PaaS (<512MB RAM)")
    ]

    for label, val in tech_list:
        p = tf_tech.add_paragraph()
        p.space_after = Pt(10)
        r1 = p.add_run()
        r1.text = "  - " + label
        r1.font.bold = True
        r1.font.size = Pt(14)
        r1.font.color.rgb = BANNER_BLUE
        r2 = p.add_run()
        r2.text = val
        r2.font.size = Pt(14)
        r2.font.color.rgb = DARK_TEXT

    # Right Box: Methodology & Pipeline Flow
    card_flow = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.6), Inches(1.2), Inches(6.2), Inches(5.6))
    card_flow.fill.solid()
    card_flow.fill.fore_color.rgb = LIGHT_BG
    card_flow.line.color.rgb = BORDER_COLOR
    tf_flow = card_flow.text_frame
    tf_flow.word_wrap = True
    tf_flow.margin_left = Inches(0.3)
    tf_flow.margin_top = Inches(0.3)

    p_fh = tf_flow.paragraphs[0]
    p_fh.text = "• Methodology & System Implementation"
    p_fh.font.size = Pt(18)
    p_fh.font.bold = True
    p_fh.font.color.rgb = NAVY
    p_fh.space_after = Pt(10)

    flow_steps = [
        "1. Weather Stream Ingestion: Open-Meteo API with SQLite LRU 1-hr caching buffer.",
        "2. Tri-Index Biomet Engine: Concurrent execution of UTCI, WBGT, and Heat Index.",
        "3. Random Forest ML HVI Engine: Fuses thermal stress with 641 district/291 ward vulnerability features.",
        "4. Decision Matrix & Simulator: Interactive GIS mapping and urban cooling mitigation sliders.",
        "5. Automated Emergency Gateway: Generates multi-tier advisories via SMS/WhatsApp."
    ]

    for fs in flow_steps:
        p = tf_flow.add_paragraph()
        p.text = fs
        p.font.size = Pt(13)
        p.font.color.rgb = DARK_TEXT
        p.space_after = Pt(8)

    # Demo Placeholders Box
    p_dh = tf_flow.add_paragraph()
    p_dh.text = "Working Prototype Demonstration:"
    p_dh.font.size = Pt(14)
    p_dh.font.bold = True
    p_dh.font.color.rgb = BANNER_BLUE
    p_dh.space_before = Pt(8)

    p_v = tf_flow.add_paragraph()
    p_v.text = "  • Video Link: [Insert YouTube Prototype Video Link Here]"
    p_v.font.size = Pt(12)
    p_v.font.color.rgb = DARK_TEXT

    p_q = tf_flow.add_paragraph()
    p_q.text = "  • Repository / Demo: [Insert GitHub Repo / QR Code Here]"
    p_q.font.size = Pt(12)
    p_q.font.color.rgb = DARK_TEXT

    # -------------------------------------------------------------
    # SLIDE 4: FEASIBILITY AND VIABILITY
    # -------------------------------------------------------------
    slide4 = prs.slides.add_slide(blank_layout)
    add_common_header_footer(slide4, "FEASIBILITY AND VIABILITY", 4)

    s4_data = [
        ("Analysis of Feasibility", [
            ("Technical Feasibility:", " Validated biometeorological algorithms running at <50ms execution latency per district."),
            ("Economic Feasibility:", " Zero proprietary software or API licensing costs; operates entirely on open-source stack and free-tier cloud."),
            ("Operational Feasibility:", " Seamlessly aligns with National Disaster Management Authority (NDMA) Heat Action Plan workflows.")
        ]),
        ("Potential Challenges & Risks", [
            ("API Rate Limits:", " Risk of external weather API throttling during severe multi-region heatwaves."),
            ("Geospatial Sparsity:", " Missing ward-level microclimate or building density metrics in tier-2/3 cities."),
            ("Server Memory Limits:", " Potential out-of-memory crashes on resource-constrained cloud hosts (<512 MB RAM).")
        ]),
        ("Mitigation Strategies", [
            ("Caching & Batching:", " 1-hour SQLite LRU cache & batch requests ensure staying under 10k daily API limits."),
            ("ML Imputation:", " Feature imputation pipeline using neighboring ward parameters when local micro-data is absent."),
            ("In-Memory Optimization:", " Vectorized GeoJSON processing without heavy GIS servers, keeping RAM usage <180 MB.")
        ])
    ]

    for idx, (title, items) in enumerate(s4_data):
        card = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.2 + idx * 1.87), Inches(12.333), Inches(1.7))
        card.fill.solid()
        card.fill.fore_color.rgb = LIGHT_BG
        card.line.color.rgb = BORDER_COLOR

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.15)

        p_t = tf.paragraphs[0]
        p_t.text = "• " + title
        p_t.font.size = Pt(17)
        p_t.font.bold = True
        p_t.font.color.rgb = NAVY
        p_t.space_after = Pt(4)

        for sub_lbl, desc in items:
            p_i = tf.add_paragraph()
            p_i.space_after = Pt(2)
            r1 = p_i.add_run()
            r1.text = "  - " + sub_lbl
            r1.font.bold = True
            r1.font.size = Pt(13)
            r1.font.color.rgb = BANNER_BLUE
            r2 = p_i.add_run()
            r2.text = desc
            r2.font.size = Pt(13)
            r2.font.color.rgb = DARK_TEXT

    # -------------------------------------------------------------
    # SLIDE 5: IMPACT AND BENEFITS
    # -------------------------------------------------------------
    slide5 = prs.slides.add_slide(blank_layout)
    add_common_header_footer(slide5, "IMPACT AND BENEFITS", 5)

    # Top Half: Target Audience Impact
    card_aud = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.2), Inches(12.333), Inches(2.6))
    card_aud.fill.solid()
    card_aud.fill.fore_color.rgb = LIGHT_BG
    card_aud.line.color.rgb = BORDER_COLOR

    tf_aud = card_aud.text_frame
    tf_aud.word_wrap = True
    tf_aud.margin_left = Inches(0.3)
    tf_aud.margin_top = Inches(0.2)

    p_ah = tf_aud.paragraphs[0]
    p_ah.text = "• Potential Impact on the Target Audience"
    p_ah.font.size = Pt(18)
    p_ah.font.bold = True
    p_ah.font.color.rgb = NAVY
    p_ah.space_after = Pt(8)

    aud_items = [
        ("Disaster Management (NDMA/SDMAs):", " Provides real-time district & ward heat risk mapping for swift cooling shelter & emergency deployment."),
        ("Healthcare & Public Safety:", " Delivers occupational WBGT heat stress alerts to protect outdoor laborers, agricultural workers & traffic police."),
        ("Urban Local Bodies (ULBs):", " Empowers urban planners with quantitative simulations for targeted green cover and cool roof investments.")
    ]

    for label, desc in aud_items:
        p = tf_aud.add_paragraph()
        p.space_after = Pt(4)
        r1 = p.add_run()
        r1.text = "  - " + label
        r1.font.bold = True
        r1.font.size = Pt(13.5)
        r1.font.color.rgb = BANNER_BLUE
        r2 = p.add_run()
        r2.text = desc
        r2.font.size = Pt(13.5)
        r2.font.color.rgb = DARK_TEXT

    # Bottom Half: Benefits (Social, Economic, Environmental)
    card_ben = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(4.0), Inches(12.333), Inches(2.8))
    card_ben.fill.solid()
    card_ben.fill.fore_color.rgb = LIGHT_BG
    card_ben.line.color.rgb = BORDER_COLOR

    tf_ben = card_ben.text_frame
    tf_ben.word_wrap = True
    tf_ben.margin_left = Inches(0.3)
    tf_ben.margin_top = Inches(0.2)

    p_bh = tf_ben.paragraphs[0]
    p_bh.text = "• Benefits of the Solution (Social, Economic, Environmental)"
    p_bh.font.size = Pt(18)
    p_bh.font.bold = True
    p_bh.font.color.rgb = NAVY
    p_bh.space_after = Pt(8)

    ben_items = [
        ("Social Benefits:", " Direct protection for vulnerable populations (laborers, elderly, urban slum communities) through hyper-local alerts."),
        ("Economic Benefits:", " Mitigates heat-induced labor productivity loss (est. 4.3% of working hours in India) and reduces emergency hospitalizations."),
        ("Environmental Benefits:", " Guides targeted cool roof and urban forestry interventions to effectively combat the Urban Heat Island (UHI) effect.")
    ]

    for label, desc in ben_items:
        p = tf_ben.add_paragraph()
        p.space_after = Pt(6)
        r1 = p.add_run()
        r1.text = "  - " + label
        r1.font.bold = True
        r1.font.size = Pt(13.5)
        r1.font.color.rgb = BANNER_BLUE
        r2 = p.add_run()
        r2.text = desc
        r2.font.size = Pt(13.5)
        r2.font.color.rgb = DARK_TEXT

    # -------------------------------------------------------------
    # SLIDE 6: RESEARCH AND REFERENCES
    # -------------------------------------------------------------
    slide6 = prs.slides.add_slide(blank_layout)
    add_common_header_footer(slide6, "RESEARCH AND REFERENCES", 6)

    card_ref = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.2), Inches(12.333), Inches(5.6))
    card_ref.fill.solid()
    card_ref.fill.fore_color.rgb = LIGHT_BG
    card_ref.line.color.rgb = BORDER_COLOR

    tf_ref = card_ref.text_frame
    tf_ref.word_wrap = True
    tf_ref.margin_left = Inches(0.4)
    tf_ref.margin_top = Inches(0.3)

    p_rh = tf_ref.paragraphs[0]
    p_rh.text = "• Details / Links of the reference and research work"
    p_rh.font.size = Pt(18)
    p_rh.font.bold = True
    p_rh.font.color.rgb = NAVY
    p_rh.space_after = Pt(12)

    refs = [
        ("1. Universal Thermal Climate Index (UTCI):", " Bröde, P., et al. (2012). 'Deriving the Operational Procedure for the Universal Thermal Climate Index (UTCI).' Int J Biometeorol 56, 481–494."),
        ("2. ISO 7243 Thermal Ergonomics Standard:", " ISO. (2017). 'ISO 7243: Ergonomics of the thermal environment - Assessment of heat stress using the WBGT index.'"),
        ("3. NOAA Heat Index Equation:", " Rothfusz, L. P. (1990). 'The Heat Index Equation.' NWS Technical Attachment SR 90-23, Scientific Services Division."),
        ("4. IMD & NDMA Action Guidelines:", " National Disaster Management Authority (NDMA), Govt of India. (2019). 'Preparation of Action Plan - Prevention & Management of Heat-Wave.'"),
        ("5. Urban Microclimate Modeling:", " Oke, T. R. (1982). 'The energetic basis of the urban heat island.' Quarterly Journal of the Royal Meteorological Society, 108(455), 1-24.")
    ]

    for title_r, desc_r in refs:
        p = tf_ref.add_paragraph()
        p.space_after = Pt(10)
        r1 = p.add_run()
        r1.text = title_r + " "
        r1.font.bold = True
        r1.font.size = Pt(14)
        r1.font.color.rgb = BANNER_BLUE
        r2 = p.add_run()
        r2.text = desc_r
        r2.font.size = Pt(14)
        r2.font.color.rgb = DARK_TEXT

    output_path = "SIH_2026_Presentation_Slides.pptx"
    prs.save(output_path)
    print(f"Presentation successfully created at {output_path}")

if __name__ == "__main__":
    create_deck()
